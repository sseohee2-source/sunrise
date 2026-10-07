"""사용료 기준(시설·기숙사·할인·공급자정보)을 엑셀 파일로 읽고 쓴다.

exe 옆의 `사용료기준.xlsx` 를 사용하며, 없으면 규정 기본값으로 새로 만든다.
규정이 바뀌면 이 엑셀만 수정하면 된다.
"""
from __future__ import annotations

import os
import sys
from dataclasses import fields

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .model import (FACILITY_GROUPS, FAMILY_DISCOUNT_KEY, NO_DISCOUNT, Discount, DormType,
                    RateBook, Room, Supplier)

RATE_FILE_NAME = "사용료기준.xlsx"

# 규모별 규정 요금 (기본 3시간, 추가 1시간당)
GROUP_FEES = {
    "강의실(50인미만)": (100_000, 20_000),
    "강의실(50인이상 100인미만)": (150_000, 30_000),
    "강의실(100인이상)": (200_000, 40_000),
    "컴퓨터실": (300_000, 20_000),
    "소강당": (300_000, 60_000),
    "대강당": (500_000, 100_000),
    "운동장": (300_000, 100_000),
}

_ROOMS = [
    # 호실, 건물, 층, 인원, 형태, 규모, 비고
    ("CEO룸", "봉사관", "1층", "20", "강의실", "강의실(50인미만)", ""),
    ("111", "봉사관", "1층", "40", "강의실", "강의실(50인미만)", "창고로 사용"),
    ("112", "봉사관", "1층", "60", "강의실", "강의실(50인이상 100인미만)", ""),
    ("118-1", "봉사관", "1층", "2", "웨비나실", "강의실(50인미만)", "연수팀 사용(대관 시 협의)"),
    ("118-2", "봉사관", "1층", "2", "웨비나실", "강의실(50인미만)", "연수팀 사용(대관 시 협의)"),
    ("118-3", "봉사관", "1층", "2", "웨비나실", "강의실(50인미만)", "연수팀 사용(대관 시 협의)"),
    ("118-4", "봉사관", "1층", "2", "웨비나실", "강의실(50인미만)", "연수지원팀"),
    ("205", "봉사관", "2층", "15", "강의실", "강의실(50인미만)", ""),
    ("207", "봉사관", "2층", "40", "강의실", "강의실(50인미만)", ""),
    ("208", "봉사관", "2층", "40", "강의실", "강의실(50인미만)", ""),
    ("209", "봉사관", "2층", "40", "강의실", "강의실(50인미만)", ""),
    ("210(PC)", "봉사관", "2층", "40", "컴퓨터실", "컴퓨터실", ""),
    ("212", "봉사관", "2층", "20", "강의실", "강의실(50인미만)", ""),
    ("301", "봉사관", "3층", "80", "강의실", "강의실(50인이상 100인미만)", ""),
    ("302", "봉사관", "3층", "70", "강의실", "강의실(50인이상 100인미만)", ""),
    ("305", "봉사관", "3층", "70", "강의실", "강의실(50인이상 100인미만)", ""),
    ("309", "봉사관", "3층", "30", "강의실", "강의실(50인미만)", ""),
    ("310", "봉사관", "3층", "40", "강의실", "강의실(50인미만)", ""),
    ("311", "봉사관", "3층", "24", "강의실", "강의실(50인미만)", "창고로 사용"),
    ("313", "봉사관", "3층", "24", "강의실", "강의실(50인미만)", ""),
    ("315", "봉사관", "3층", "24", "강의실", "강의실(50인미만)", ""),
    ("317", "봉사관", "3층", "24", "강의실", "강의실(50인미만)", ""),
    ("322", "봉사관", "3층", "30", "강의실", "강의실(50인미만)", ""),
    ("326", "봉사관", "3층", "30", "강의실", "강의실(50인미만)", ""),
    ("327", "봉사관", "3층", "40", "강의실", "강의실(50인미만)", ""),
    ("328", "봉사관", "3층", "40", "강의실", "강의실(50인미만)", ""),
    ("다목적강의실", "인화2관", "1층", "100", "강의실", "강의실(100인이상)", ""),
    ("비즈니스룸", "인화2관", "1층", "10", "강의실", "강의실(50인미만)", ""),
    ("운동장", "운동장", "야외", "500~1,000", "운동장", "운동장", ""),
    ("소강당", "봉사관", "1층", "162", "소강당", "소강당", ""),
    ("대강당", "창의관", "1층", "421", "대강당", "대강당", ""),
    ("실험동 101", "실험동", "1층", "30", "강의실", "강의실(50인미만)", "대관X, 필요시 협의"),
    ("실험동 106", "실험동", "1층", "30", "강의실", "강의실(50인미만)", "대관X, 필요시 협의"),
    ("실험동 201", "실험동", "2층", "30", "강의실", "강의실(50인미만)", "대관X, 필요시 협의"),
    ("실험동 209", "실험동", "2층", "30", "강의실", "강의실(50인미만)", "대관X, 필요시 협의"),
    ("실험동 210", "실험동", "2층", "30", "강의실", "강의실(50인미만)", "대관X, 필요시 협의"),
]

_DORMS = [("1인실", 30_000), ("2인실", 40_000), ("콘도형 A", 60_000), ("단체실 및 콘도형 B", 100_000)]

_DISCOUNTS = [
    ("[면제]정부 부처 및 지방자치단체", "면 제", 1.0),
    ("[면제]중진공과 공동으로 사업을 추진하는 공공기관 또는 유관단체", "면 제", 1.0),
    ("[면제]연수원 소재 지역 관공서 및 상공회의소 등 연수원과 유대가 필요한 지역 사회단체", "면 제", 1.0),
    ("[면제]협약기간내의 청년창업사관학교 입교생이 사용하는 창업준비공간(장비 포함)", "면 제", 1.0),
    ("[면제]그 밖의 시설개방 목적상 사용료를 면제함이 타당하다고 인정되는 자", "면 제", 1.0),
    ("[50%감면]공공기관 및 유관단체", "50%감면", 0.5),
    ("[50%감면](사)글로벌최고경영자클럽, (사)중소기업융합중앙회 등 유관 중소기업 경영자 협·단체 및 연수수료자 동우회",
     "50%감면", 0.5),
    ("[50%감면]협약기간내의 청년창업사관학교 입교생 및 경영후계자", "50%감면", 0.5),
    ("[50%감면]그 밖에 시설개방 목적상 사용료를 감면함이 타당하다고 인정되는 자", "50%감면", 0.5),
    ("[할인] 패밀리기업 할인", "할인", None),
]


def default_book() -> RateBook:
    rooms = []
    for name, bld, fl, cap, kind, group, note in _ROOMS:
        base, extra = GROUP_FEES[group]
        rooms.append(Room(name, bld, fl, cap, kind, group, base, extra, available=not note, note=note))
    return RateBook(rooms=rooms,
                    dorms=[DormType(n, f) for n, f in _DORMS],
                    discounts=[Discount(n, c, r) for n, c, r in _DISCOUNTS],
                    supplier=Supplier())


# ---------------------------------------------------------------- 파일 위치
def app_dir() -> str:
    """exe(또는 실행 스크립트)가 있는 폴더."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def resource_path(*parts: str) -> str:
    """exe 내부에 포함된 리소스 경로."""
    base = getattr(sys, "_MEIPASS", None)
    base = os.path.join(base, "quote_app") if base else os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, *parts)


def rate_file_path() -> str:
    return os.path.join(app_dir(), RATE_FILE_NAME)


# ---------------------------------------------------------------- 읽기/쓰기
ROOM_HEADERS = ["호실", "건물명", "층", "수용인원", "형태", "규모(구분)", "기본(3시간)", "추가 1시간당",
                "대관가능(O/X)", "비고"]
DORM_HEADERS = ["구분", "1박 사용료"]
DISC_HEADERS = ["대상", "구분", "할인률(%)", "비고"]
SUPPLIER_LABELS = {
    "title": "견적명", "name": "공급자", "biz_no": "사업자번호", "ceo": "대표", "biz_type": "업태",
    "biz_item": "종목", "address": "주소", "manager": "담당자", "manager_email": "이메일",
    "office_tel": "사무실", "fax": "FAX", "cafeteria_note": "구내식당 안내문",
    "footer_note": "견적서 하단 안내문", "notice": "유의사항", "family_guide": "패밀리기업 할인 안내",
}

_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="404040")


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(1, i)
        c.font, c.fill = _HEAD_FONT, _HEAD_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def save_book(book: RateBook, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "시설"
    _header(ws, ROOM_HEADERS, [14, 10, 7, 11, 10, 28, 12, 13, 13, 28])
    for r in book.rooms:
        ws.append([r.name, r.building, r.floor, r.capacity, r.kind, r.group, r.base_fee, r.extra_fee,
                   "O" if r.available else "X", r.note])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in row:
            c.number_format = "#,##0"
    dv = DataValidation(type="list", formula1='"' + ",".join(FACILITY_GROUPS) + '"', allow_blank=False)
    dv.add("F2:F500")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"O,X"')
    dv2.add("I2:I500")
    ws.add_data_validation(dv2)

    ws = wb.create_sheet("기숙사")
    _header(ws, DORM_HEADERS, [22, 14])
    for d in book.dorms:
        ws.append([d.name, d.fee])
        ws.cell(ws.max_row, 2).number_format = "#,##0"

    ws = wb.create_sheet("할인기준")
    _header(ws, DISC_HEADERS, [90, 10, 11, 40])
    for d in book.discounts:
        ws.append([d.name, d.category, None if d.rate is None else round(d.rate * 100, 4),
                   "비워두면 담당자가 행별로 할인률 선택" if d.rate is None else ""])

    ws = wb.create_sheet("공급자정보")
    _header(ws, ["항목", "값", "키(수정금지)"], [20, 90, 16])
    for f in fields(Supplier):
        ws.append([SUPPLIER_LABELS.get(f.name, f.name), getattr(book.supplier, f.name), f.name])
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _i(v) -> int:
    if v in (None, ""):
        return 0
    return int(round(float(str(v).replace(",", ""))))


def load_book(path: str) -> RateBook:
    wb = load_workbook(path, data_only=True)
    rooms = []
    for row in wb["시설"].iter_rows(min_row=2, values_only=True):
        if not row or not _s(row[0]):
            continue
        row = list(row) + [None] * 10
        rooms.append(Room(_s(row[0]), _s(row[1]), _s(row[2]), _s(row[3]), _s(row[4]), _s(row[5]),
                          _i(row[6]), _i(row[7]), available=_s(row[8]).upper() != "X", note=_s(row[9])))
    dorms = [DormType(_s(r[0]), _i(r[1])) for r in wb["기숙사"].iter_rows(min_row=2, values_only=True)
             if r and _s(r[0])]
    discounts = []
    for r in wb["할인기준"].iter_rows(min_row=2, values_only=True):
        if not r or not _s(r[0]):
            continue
        rate = r[2] if len(r) > 2 else None
        rate = None if rate in (None, "") else float(str(rate).replace("%", "")) / 100
        discounts.append(Discount(_s(r[0]), _s(r[1]), rate))
    sup = Supplier()
    if "공급자정보" in wb.sheetnames:
        names = {f.name for f in fields(Supplier)}
        for r in wb["공급자정보"].iter_rows(min_row=2, values_only=True):
            if r and len(r) > 2 and r[2] in names:
                setattr(sup, r[2], "" if r[1] is None else str(r[1]))
    return RateBook(rooms, dorms, discounts, sup)


def load_or_create(path: str | None = None) -> tuple[RateBook, str]:
    path = path or rate_file_path()
    if not os.path.exists(path):
        book = default_book()
        try:
            save_book(book, path)
        except OSError:
            pass
        return book, path
    return load_book(path), path


def discount_choices(book: RateBook) -> list[Discount]:
    return [Discount(NO_DISCOUNT, "해당없음", 0.0)] + book.discounts


def is_family(d: Discount | None) -> bool:
    return d is not None and (d.rate is None or FAMILY_DISCOUNT_KEY in d.name)
